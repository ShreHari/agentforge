"""Turn an uploaded CV file (PDF or Word) into plain text.

This is plain code, not AI: the agent only ever sees text. The user checks and edits the
text before a run starts, so a jumbled two-column PDF can be fixed by the person who
knows what it should say, and they see exactly what the AI will read.

Uploaded files are untrusted, so every limit below is checked BEFORE the heavy work.
"""

import io
import zipfile

import pdfplumber
from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_BYTES = 2 * 1024 * 1024  # 2 MB. A CV is rarely over 500 KB
MAX_PAGES = 10
MAX_UNZIPPED_BYTES = 20 * 1024 * 1024  # a .docx is a zip; this stops "zip bombs"
X_TOLERANCE = 1.5  # see pdf_text() for how this number was chosen

# We check what a file IS by its first bytes, not by its name or the type the browser
# claims: anyone can rename virus.exe to cv.pdf.
PDF_START = b"%PDF-"  # every PDF starts with these bytes
ZIP_START = b"PK\x03\x04"  # every zip (and so every .docx) starts with these


class ParseError(ValueError):
    """The file can't be turned into CV text. The message is safe to show the user."""


def extract_text(data: bytes) -> str:
    """Plain text from a PDF or .docx file, tidied. Raises ParseError if it can't."""
    if len(data) > MAX_BYTES:
        raise ParseError("The file is too big. The limit is 2 MB.")
    if data.startswith(PDF_START):
        text = pdf_text(data)
    elif data.startswith(ZIP_START):
        text = docx_text(data)
    else:
        raise ParseError("Only PDF and Word (.docx) files are supported.")

    text = tidy(text)
    if not text:
        # A scanned CV is a picture of text, so there's no text to extract.
        raise ParseError("No text found. If this is a scanned PDF, please paste your CV as text.")
    return text


def pdf_text(data: bytes) -> str:
    # Step 1, the quick safety checks, with pypdf: it reads the PDF's structure fast.
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ParseError("This PDF is password-protected. Please upload one without a password.")
        if len(reader.pages) > MAX_PAGES:
            raise ParseError(f"The PDF has over {MAX_PAGES} pages. Is this really a CV?")
    except PdfReadError as error:
        raise ParseError("This PDF looks damaged. Try saving it again, or paste the text.") from error

    # Step 2, the text, with pdfplumber. A PDF often stores no space characters, only
    # where to draw each letter, so the library must GUESS where spaces go from the gaps.
    # x_tolerance is how close (in points) two letters must be to count as the same word.
    # Measured on a real CV (25 Oct, practice/compare_pdf_text.py): words glued together
    # pypdf 64, pdfplumber default (3) 72, pdfplumber 1.5: 0, with no words split apart.
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join(page.extract_text(x_tolerance=X_TOLERANCE) or "" for page in pdf.pages)
    except Exception as error:  # pdfplumber fails in several ways on a broken PDF
        raise ParseError("This PDF looks damaged. Try saving it again, or paste the text.") from error


def docx_text(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            # Add up the sizes the files WOULD have when unzipped, before unzipping anything.
            if sum(item.file_size for item in archive.infolist()) > MAX_UNZIPPED_BYTES:
                raise ParseError("This Word file is far too large inside. Please paste the text.")
        doc = Document(io.BytesIO(data))
    except ParseError:
        raise
    except Exception as error:  # a broken or non-Word zip fails in several different ways
        raise ParseError("This doesn't look like a Word (.docx) file, or it is damaged.") from error

    parts = [paragraph.text for paragraph in doc.paragraphs]
    # Many CV templates put skills or dates in tables, which paragraphs don't include.
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def tidy(text: str) -> str:
    """Squash runs of spaces, and keep at most one blank line in a row."""
    lines = [" ".join(line.split()) for line in text.splitlines()]
    kept: list[str] = []
    for line in lines:
        if line or (kept and kept[-1]):  # skip a blank line if the last one was blank too
            kept.append(line)
    return "\n".join(kept).strip()
