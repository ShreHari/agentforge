"""CV upload: PDF and Word files become plain text. No AI, no database.

Each test builds a small file in memory, so there are no binary files to keep in git.
"""

import io
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from api.main import app
from api.parse import MAX_BYTES, ParseError, extract_text, tidy


def make_pdf(lines: list[str]) -> bytes:
    """A real one-page PDF containing these lines of text."""
    text = " ".join(f"({line}) Tj 0 -16 Td" for line in lines)
    return pdf_from_content(f"BT /F1 12 Tf 72 720 Td {text} ET".encode())


def pdf_from_content(content: bytes) -> bytes:
    """A one-page PDF that draws `content` (PDF drawing commands) with Helvetica.

    A PDF is numbered "objects" plus a table (xref) saying where each one starts.
    Object 4 is the page's content.
    """
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
         b"/Resources << /Font << /F1 5 0 R >> >> >>"),
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1, xref_at)
    return bytes(out)


def blank_pdf(pages: int = 1, password: str | None = None) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=612, height=792)
    if password:
        writer.encrypt(password)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("Alex Example")
    doc.add_paragraph("Built REST APIs with FastAPI")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Python, Docker"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# --- extract_text ---


def test_pdf_text_is_extracted():
    text = extract_text(make_pdf(["Alex Example", "Built REST APIs with FastAPI"]))
    assert "Alex Example" in text
    assert "Built REST APIs with FastAPI" in text


def test_words_drawn_with_gaps_not_spaces_keep_their_spaces():
    """Many CV exports draw the gap between words as a small nudge (-150 below), with no
    space character. pdfplumber's default setting glues these into "Built6+FastAPI";
    our X_TOLERANCE doesn't. If someone changes the setting, this test catches it."""
    content = b"BT /F1 11 Tf 72 720 Td [(Built) -150 (6+) -150 (FastAPI) -150 (endpoints)] TJ ET"
    assert extract_text(pdf_from_content(content)) == "Built 6+ FastAPI endpoints"


def test_docx_paragraphs_and_tables_are_extracted():
    text = extract_text(make_docx())
    assert "Built REST APIs with FastAPI" in text
    assert "Skills | Python, Docker" in text  # tables are often where skills live


def test_scanned_pdf_with_no_text_is_refused():
    with pytest.raises(ParseError, match="scanned"):
        extract_text(blank_pdf())


def test_password_protected_pdf_is_refused():
    with pytest.raises(ParseError, match="password"):
        extract_text(blank_pdf(password="secret"))


def test_too_many_pages_is_refused():
    with pytest.raises(ParseError, match="pages"):
        extract_text(blank_pdf(pages=11))


def test_zip_bomb_is_refused_before_unzipping():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (21 * 1024 * 1024))  # 21 MB unzipped
    data = buffer.getvalue()
    assert len(data) < MAX_BYTES  # tiny when zipped, so the size limit alone won't catch it
    with pytest.raises(ParseError, match="too large inside"):
        extract_text(data)


def test_a_renamed_file_is_judged_by_its_bytes_not_its_name():
    with pytest.raises(ParseError, match="Only PDF and Word"):
        extract_text(b"MZ this is really a program, renamed to cv.pdf")


def test_tidy_squashes_spaces_and_blank_lines():
    assert tidy("  Alex    Example  \n\n\n\nPython   \n") == "Alex Example\n\nPython"


# --- the /parse endpoint ---


@pytest.fixture
def client():
    with TestClient(app) as test_client:
        yield test_client


def test_parse_endpoint_returns_text(client):
    response = client.post(
        "/parse", files={"file": ("cv.pdf", make_pdf(["Alex Example"]), "application/pdf")})
    assert response.status_code == 200
    body = response.json()
    assert "Alex Example" in body["text"]
    assert body["characters"] == len(body["text"])
    assert body["too_long"] is False


def test_parse_endpoint_explains_a_bad_file(client):
    response = client.post("/parse", files={"file": ("cv.txt", b"hello", "text/plain")})
    assert response.status_code == 422
    assert "Only PDF and Word" in response.json()["detail"]


def test_parse_endpoint_refuses_a_huge_file(client):
    huge = b"%PDF-" + b"0" * MAX_BYTES
    response = client.post("/parse", files={"file": ("cv.pdf", huge, "application/pdf")})
    assert response.status_code == 413
